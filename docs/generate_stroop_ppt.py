from __future__ import annotations

import random
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


COLORS = {
    "RED": RGBColor(192, 0, 0),
    "BLUE": RGBColor(0, 102, 204),
    "GREEN": RGBColor(0, 153, 0),
    "YELLOW": RGBColor(204, 153, 0),
    "PINK": RGBColor(204, 0, 102),
    "PURPLE": RGBColor(102, 0, 153),
    "BLACK": RGBColor(0, 0, 0),
    "ORANGE": RGBColor(230, 120, 0),
}
WORDS = list(COLORS)
MISMATCHED_COLOR_CHOICES = {
    word: [name for name in COLORS if name != word]
    for word in WORDS
}

LAYOUTS = [
    (6, 6),
    (5, 7),
    (4, 8),
    (7, 5),
    (3, 10),
    (8, 4),
]

SLIDE_WIDTH_INCHES = 13.333
SLIDE_HEIGHT_INCHES = 7.5
LEFT_MARGIN_INCHES = 0.7
TOP_MARGIN_INCHES = 0.75
USABLE_WIDTH_INCHES = 11.9
USABLE_HEIGHT_INCHES = 6.1
TITLE_LEFT_INCHES = 0.5
TITLE_TOP_INCHES = 0.18
TITLE_WIDTH_INCHES = 5.5
TITLE_HEIGHT_INCHES = 0.4
FOOTER_LEFT_INCHES = 11.9
FOOTER_TOP_INCHES = 7.0
FOOTER_WIDTH_INCHES = 0.9
FOOTER_HEIGHT_INCHES = 0.25


def build_deck() -> Presentation:
    rng = random.Random(37048)
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_INCHES)
    presentation.slide_height = Inches(SLIDE_HEIGHT_INCHES)

    blank = presentation.slide_layouts[6]

    for slide_number in range(30):
        rows, columns = LAYOUTS[slide_number % len(LAYOUTS)]
        slide = presentation.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

        slide.shapes.add_textbox(
            Inches(TITLE_LEFT_INCHES),
            Inches(TITLE_TOP_INCHES),
            Inches(TITLE_WIDTH_INCHES),
            Inches(TITLE_HEIGHT_INCHES),
        )
        title_frame = slide.shapes[-1].text_frame
        title_frame.text = f"STROOP TEST {slide_number + 1:02d}"
        title_run = title_frame.paragraphs[0].runs[0]
        title_run.font.bold = True
        title_run.font.size = Pt(18)
        title_run.font.color.rgb = RGBColor(90, 90, 90)

        cell_width = USABLE_WIDTH_INCHES / columns
        cell_height = USABLE_HEIGHT_INCHES / rows

        for row in range(rows):
            for column in range(columns):
                word = WORDS[(row * columns + column + slide_number) % len(WORDS)]
                color_choices = MISMATCHED_COLOR_CHOICES[word]
                font_color_name = color_choices[rng.randrange(len(color_choices))]

                x = Inches(LEFT_MARGIN_INCHES + column * cell_width)
                y = Inches(TOP_MARGIN_INCHES + row * cell_height)
                width = Inches(cell_width)
                height = Inches(cell_height)

                textbox = slide.shapes.add_textbox(x, y, width, height)
                frame = textbox.text_frame
                frame.clear()
                paragraph = frame.paragraphs[0]
                paragraph.alignment = PP_ALIGN.CENTER
                run = paragraph.add_run()
                run.text = word
                run.font.bold = True
                run.font.size = Pt(26 if columns <= 6 else 22)
                run.font.name = "Arial"
                run.font.color.rgb = COLORS[font_color_name]

        footer = slide.shapes.add_textbox(
            Inches(FOOTER_LEFT_INCHES),
            Inches(FOOTER_TOP_INCHES),
            Inches(FOOTER_WIDTH_INCHES),
            Inches(FOOTER_HEIGHT_INCHES),
        )
        footer_frame = footer.text_frame
        footer_frame.text = str(slide_number + 1)
        footer_paragraph = footer_frame.paragraphs[0]
        footer_paragraph.alignment = PP_ALIGN.RIGHT
        footer_run = footer_paragraph.runs[0]
        footer_run.font.size = Pt(11)
        footer_run.font.color.rgb = RGBColor(120, 120, 120)

    return presentation


def main() -> None:
    output_path = Path(__file__).with_name("stroop-test-30-slides.pptx")
    presentation = build_deck()
    presentation.save(output_path)
    print(f"Generated {output_path}")
    print(f"Slides: {len(presentation.slides)}")


if __name__ == "__main__":
    main()
